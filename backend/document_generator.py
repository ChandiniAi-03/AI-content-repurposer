from docx import Document


def create_document(twitter_content, linkedin_content, instagram_content):

    document = Document()

    # Title
    document.add_heading("AI Content Repurposer", level=0)

    document.add_paragraph(
        "Generated social-media content from an article"
    )

    # Twitter/X
    document.add_heading("🐦 Twitter / X Thread", level=1)
    document.add_paragraph(twitter_content)

    # LinkedIn
    document.add_heading("💼 LinkedIn Post", level=1)
    document.add_paragraph(linkedin_content)

    # Instagram
    document.add_heading("📸 Instagram Caption", level=1)
    document.add_paragraph(instagram_content)

    return document